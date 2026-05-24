"""
PDF Template endpoints.

GET    /pdf-templates                    → list org's templates
GET    /pdf-templates/base               → download base DOCX scaffold
POST   /pdf-templates                    → upload DOCX template (multipart)
PATCH  /pdf-templates/{id}              → rename
POST   /pdf-templates/{id}/set-default  → set as org default
DELETE /pdf-templates/{id}              → delete
"""
import io

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.database import get_db
from app.core.permissions import require_permission
from app.schemas.pdf_template import PdfTemplateOut, PdfTemplateRename
from app.services import pdf_template_service

router = APIRouter(tags=["pdf_templates"])


def _build_base_docx() -> bytes:
    """Build a base DOCX scaffold with the {{IRDoc_CONTENT}} marker."""
    doc = Document()

    # ── Cover page ────────────────────────────────────────────────
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("INCIDENT REPORT")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle.add_run("IRDoc — Incident Response Platform")
    sub_run.font.size = Pt(12)
    sub_run.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    doc.add_paragraph()  # spacer

    instructions = doc.add_paragraph()
    instructions.alignment = WD_ALIGN_PARAGRAPH.CENTER
    instr_run = instructions.add_run(
        "Customise this document in Word (logo, fonts, colours, page layout), "
        "then upload it in Admin → Report Templates.\n"
        "Do not remove or rename the {{IRDoc_CONTENT}} marker below."
    )
    instr_run.font.size = Pt(10)
    instr_run.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)
    instr_run.italic = True

    doc.add_page_break()

    # ── Content marker ────────────────────────────────────────────
    marker_para = doc.add_paragraph("{{IRDoc_CONTENT}}")
    marker_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

_MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB

_DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


@router.get("/pdf-templates/base")
async def download_base_template(
    current_user=Depends(require_permission("reports.read")),
):
    docx_bytes = _build_base_docx()
    return StreamingResponse(
        io.BytesIO(docx_bytes),
        media_type=_DOCX_MIME,
        headers={"Content-Disposition": 'attachment; filename="irdoc_base_template.docx"'},
    )


@router.get("/pdf-templates")
async def list_pdf_templates(
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_permission("reports.read")),
):
    templates = await pdf_template_service.list_templates(str(current_user.org_id), db)
    return {
        "data": [PdfTemplateOut.model_validate(t) for t in templates],
        "meta": {"total": len(templates)},
        "error": None,
    }


@router.post("/pdf-templates", status_code=201)
async def upload_pdf_template(
    name: str = Form(...),
    file: UploadFile = File(...),
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_permission("templates.create")),
):
    if not file.filename or not file.filename.lower().endswith(".docx"):
        raise HTTPException(status_code=400, detail="Only .docx files are accepted")

    file_bytes = await file.read()
    if len(file_bytes) > _MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="File too large (max 20 MB)")

    try:
        Document(io.BytesIO(file_bytes))
    except Exception:
        raise HTTPException(status_code=400, detail="File is not a valid .docx document")

    try:
        t = await pdf_template_service.create_template(
            org_id=str(current_user.org_id),
            name=name,
            file_bytes=file_bytes,
            file_size=len(file_bytes),
            created_by=str(current_user.id),
            db=db,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc))

    return {"data": PdfTemplateOut.model_validate(t), "error": None}


@router.patch("/pdf-templates/{template_id}")
async def rename_pdf_template(
    template_id: str,
    body: PdfTemplateRename,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_permission("templates.create")),
):
    t = await pdf_template_service.rename_template(
        template_id, str(current_user.org_id), body.name, db
    )
    return {"data": PdfTemplateOut.model_validate(t), "error": None}


@router.post("/pdf-templates/{template_id}/set-default")
async def set_default_pdf_template(
    template_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_permission("templates.create")),
):
    t = await pdf_template_service.set_default(
        template_id, str(current_user.org_id), db
    )
    return {"data": PdfTemplateOut.model_validate(t), "error": None}


@router.delete("/pdf-templates/{template_id}", status_code=204)
async def delete_pdf_template(
    template_id: str,
    db: AsyncSession = Depends(get_db),
    current_user=Depends(require_permission("templates.delete")),
):
    await pdf_template_service.delete_template(
        template_id, str(current_user.org_id), db
    )
